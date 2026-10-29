# -*- coding: utf-8 -*-
"""野村證券資料の要約デッキ用ヘルパー（FreeForm_Form16-10(メイリオ) テンプレート専用）

使い方:
    from nomura_deck import Deck, C
    d = Deck("FreeForm_Form16-10(メイリオ).pptx")
    d.cover("タイトル", ["サブ1", "サブ2"])
    s = d.slide("スライドタイトル（結論を1文で）")
    d.card(s, x, y, w, h, "見出し", ["本文1", "本文2"])
    d.save("out.pptx")

座標はインチ。スライドは 10 x 6.25 in（16:10）。
本文エリアの目安: x 0.4〜9.6 / y 0.95〜5.35（右下にスライド番号）
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree

FONT = "メイリオ"


class C:
    """配色ルール: 上昇・引き上げ・追い風 = RED / 下落・引き下げ・逆風 = NAVY / 中立 = GRAY"""
    RED = "C00000"      # テンプレートの箇条書き記号と同じ濃い赤（強調・上昇）
    NAVY = "1F497D"     # テーマの dk2（下落・引き下げ）
    TEXT = "262626"
    SUB = "595959"
    GRAY = "808080"
    LINE = "BFBFBF"
    BG = "F2F2F2"       # カード地
    RED_T = "F9E5E5"    # 赤の淡色
    NAVY_T = "E4EBF4"   # 紺の淡色
    WHITE = "FFFFFF"
    # CIO VIEW 5段階（UW → OW）
    SCALE = ["1F497D", "7A93B8", "A6A6A6", "E08A8A", "C00000"]
    SCALE_LABEL = ["アンダー\nウェイト", "やや\nアンダー", "ニュートラル", "やや\nオーバー", "オーバー\nウェイト"]


def _rgb(h):
    return RGBColor.from_string(h)


class Deck:
    def __init__(self, template):
        self.prs = Presentation(template)
        lst = self.prs.slides._sldIdLst
        for sid in list(lst):  # テンプレートの空スライドを除去
            self.prs.part.drop_rel(sid.rId)
            lst.remove(sid)
        self.L = {l.name: l for l in self.prs.slide_layouts}
        self.W = self.prs.slide_width / 914400
        self.H = self.prs.slide_height / 914400

    # ---------- スライド ----------
    def _add_slide_number(self, slide, layout):
        for ph in layout.placeholders:
            if ph.placeholder_format.type is not None and "SLIDE_NUMBER" in str(ph.placeholder_format.type):
                el = copy.deepcopy(ph._element)
                slide.shapes._spTree.append(el)
                # テキストをフィールドに置換
                txBody = el.find(qn("p:txBody"))
                for p in txBody.findall(qn("a:p")):
                    txBody.remove(p)
                p = etree.SubElement(txBody, qn("a:p"))
                fld = etree.SubElement(p, qn("a:fld"), id="{D2D8002D-B5B0-4BAC-B1F6-782DDCCE6D9C}", type="slidenum")
                etree.SubElement(fld, qn("a:rPr"), lang="ja-JP")
                t = etree.SubElement(fld, qn("a:t"))
                t.text = "‹#›"
                # 位置はレイアウトから継承させる
                spPr = el.find(qn("p:spPr"))
                for ch in list(spPr):
                    spPr.remove(ch)
                return

    def cover(self, title, sub_lines, photo=None, photo_w=4.4):
        """表紙。photo を渡すと右側に写真を全面配置し、タイトルを左寄せにする。"""
        layout = self.L["タイトル スライド"]
        s = self.prs.slides.add_slide(layout)
        title_ph = s.shapes.title
        sub = [ph for ph in s.placeholders if ph.placeholder_format.idx == 1][0]
        if photo:
            self.image_fill(s, photo, self.W - photo_w, 0, photo_w, self.H)
            tw = self.W - photo_w - 0.9
            for ph, (y, h) in ((title_ph, (1.5, 1.7)), (sub, (3.35, 1.6))):
                ph.left, ph.top, ph.width, ph.height = Inches(0.5), Inches(y), Inches(tw), Inches(h)
            # 写真を最背面へ
            sp = s.shapes._spTree
            pic = sp[-1]
            sp.remove(pic)
            sp.insert(2, pic)
        title_ph.text = title
        for p in title_ph.text_frame.paragraphs:
            if photo:
                p.alignment = PP_ALIGN.LEFT
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(30)
        tf = sub.text_frame
        for i, line in enumerate(sub_lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            if photo:
                p.alignment = PP_ALIGN.LEFT
            p.space_after = Pt(4)
            r = p.add_run()
            r.text = line
            r.font.size = Pt(15 if i == 0 else 11)
            r.font.bold = i == 0
            r.font.color.rgb = _rgb(C.RED if i == 0 else C.SUB)
        return s

    # ---------- 画像 ----------
    def _prep(self, path, box_w=None, box_h=None, fill=False):
        """画像をJPEG化（必要ならボックス比率でトリミング）して一時ファイルを返す"""
        import os
        import tempfile
        from PIL import Image
        im = Image.open(path).convert("RGB")
        if fill and box_w and box_h:
            r = box_w / box_h
            w, h = im.size
            if w / h > r:
                nw = int(h * r)
                im = im.crop(((w - nw) // 2, 0, (w - nw) // 2 + nw, h))
            else:
                nh = int(w / r)
                im = im.crop((0, (h - nh) // 2, w, (h - nh) // 2 + nh))
        fd, out = tempfile.mkstemp(suffix=".jpg")
        os.close(fd)
        im.save(out, "JPEG", quality=88)
        return out, im.size

    def image_fill(self, slide, path, x, y, w, h, name=None):
        """写真をボックスいっぱいに配置（はみ出す分は中央基準でトリミング）"""
        f, _ = self._prep(path, w, h, fill=True)
        pic = slide.shapes.add_picture(f, Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            pic.name = name
        return pic

    def image_fit(self, slide, path, x, y, w, h, caption=None, align="c", name=None):
        """原資料の図表をボックス内に縦横比を保って配置。caption は下に小さく表示（出典ページなど）。"""
        cap_h = 0.22 if caption else 0
        f, (iw, ih) = self._prep(path)
        bw, bh = w, h - cap_h
        sc = min(bw / iw, bh / ih)
        pw, ph = iw * sc, ih * sc
        px = x + (bw - pw) / 2 if align == "c" else x
        py = y + (bh - ph) / 2
        pic = slide.shapes.add_picture(f, Inches(px), Inches(py), Inches(pw), Inches(ph))
        pic.name = name or "原資料の図表"
        pic.line.color.rgb = _rgb("D9D9D9")
        pic.line.width = Pt(0.75)
        if caption:
            self.text(slide, px, py + ph + 0.02, max(pw, 2.0), 0.22, caption, size=9, color=C.GRAY)
        return pic, (px, py, pw, ph)

    def slide(self, title, size=24):
        """「タイトルのみ」レイアウト。本文は自由配置する。
        ※テンプレートからEIZOロゴは削除済み（2026-10-29）。"""
        layout = self.L["タイトルのみ"]
        s = self.prs.slides.add_slide(layout)
        for ph in list(s.placeholders):
            if ph.placeholder_format.idx != 0:
                ph._element.getparent().remove(ph._element)
        s.shapes.title.text = title
        tf = s.shapes.title.text_frame
        for r in tf.paragraphs[0].runs:
            r.font.size = Pt(size)
            r.font.bold = True
        self._add_slide_number(s, layout)
        return s

    def notes(self, slide, text):
        slide.notes_slide.notes_text_frame.text = text

    # ---------- 基本図形 ----------
    @staticmethod
    def _font(run, size, color, bold=False, italic=False):
        f = run.font
        f.name = FONT
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.color.rgb = _rgb(color)
        rPr = run._r.get_or_add_rPr()
        for tag in ("a:ea",):
            e = rPr.find(qn(tag))
            if e is None:
                e = etree.SubElement(rPr, qn(tag))
            e.set("typeface", FONT)

    def text(self, slide, x, y, w, h, paras, size=12, color=C.TEXT, bold=False, align="l",
             anchor="t", bullet=False, space_after=4, margin=0.04, name=None, line_spacing=None):
        """paras: str | [str | [(text, {size,color,bold})...]]"""
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        m = Inches(margin)
        tf.margin_left = tf.margin_right = m
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        self._fill_tf(tf, paras, size, color, bold, align, bullet, space_after, line_spacing)
        return tb

    def _fill_tf(self, tf, paras, size, color, bold, align, bullet, space_after, line_spacing=None):
        if isinstance(paras, str):
            paras = [paras]
        al = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = al
            p.space_after = Pt(space_after)
            if line_spacing:
                p.line_spacing = line_spacing
            runs = para if isinstance(para, list) else [(para, {})]
            for txt, o in runs:
                r = p.add_run()
                r.text = txt
                self._font(r, o.get("size", size), o.get("color", color), o.get("bold", bold))
            if bullet:
                pPr = p._p.get_or_add_pPr()
                pPr.set("marL", str(int(Inches(0.17))))
                pPr.set("indent", str(-int(Inches(0.17))))
                buClr = etree.SubElement(pPr, qn("a:buClr"))
                etree.SubElement(buClr, qn("a:srgbClr"), val=C.RED)
                etree.SubElement(pPr, qn("a:buSzPct"), val="80000")
                etree.SubElement(pPr, qn("a:buFont"), typeface="Wingdings", pitchFamily="2", charset="2")
                etree.SubElement(pPr, qn("a:buChar"), char="n")

    def box(self, slide, x, y, w, h, fill=C.BG, line=None, shape="rrect", radius=0.08, name=None,
            text=None, size=12, color=C.TEXT, bold=False, align="c", anchor="m", line_w=1.0, dash=False):
        st = {"rrect": MSO_SHAPE.ROUNDED_RECTANGLE, "rect": MSO_SHAPE.RECTANGLE, "oval": MSO_SHAPE.OVAL,
              "chevron": MSO_SHAPE.CHEVRON, "pent": MSO_SHAPE.PENTAGON, "rarrow": MSO_SHAPE.RIGHT_ARROW,
              "darrow": MSO_SHAPE.DOWN_ARROW, "tri": MSO_SHAPE.ISOSCELES_TRIANGLE}[shape]
        sp = slide.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            sp.name = name
        if shape == "rrect":
            sp.adjustments[0] = min(0.5, radius / max(0.01, min(w, h)))
        if fill:
            sp.fill.solid()
            sp.fill.fore_color.rgb = _rgb(fill)
        else:
            sp.fill.background()
        if line:
            sp.line.color.rgb = _rgb(line)
            sp.line.width = Pt(line_w)
            if dash:
                from pptx.enum.dml import MSO_LINE_DASH_STYLE
                sp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            sp.line.fill.background()
        sp.shadow.inherit = False
        st_el = sp._element.find(qn("p:style"))
        if st_el is not None:  # テーマの影効果を参照しない
            er = st_el.find(qn("a:effectRef"))
            if er is not None:
                er.set("idx", "0")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        tf.word_wrap = True
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        if text is not None:
            self._fill_tf(tf, text, size, color, bold, align, False, 0)
        return sp

    def arrow(self, slide, x1, y1, x2, y2, color=C.GRAY, width=1.75, dash=False):
        ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        ln.line.color.rgb = _rgb(color)
        ln.line.width = Pt(width)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        lnEl = ln.line._get_or_add_ln()
        etree.SubElement(lnEl, qn("a:tailEnd"), type="triangle", w="med", len="med")
        return ln

    def line(self, slide, x1, y1, x2, y2, color=C.LINE, width=1.0):
        ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        ln.line.color.rgb = _rgb(color)
        ln.line.width = Pt(width)
        return ln

    # ---------- 部品 ----------
    def num_badge(self, slide, x, y, n, color=C.RED, d=0.36, size=13):
        return self.box(slide, x, y, d, d, fill=color, shape="oval", text=str(n), size=size,
                        color=C.WHITE, bold=True)

    def tag(self, slide, x, y, w, label, color=C.RED, h=0.26, size=10, filled=True):
        if filled:
            return self.box(slide, x, y, w, h, fill=color, text=label, size=size, color=C.WHITE, bold=True,
                            radius=0.13)
        return self.box(slide, x, y, w, h, fill=C.WHITE, line=color, text=label, size=size, color=color,
                        bold=True, radius=0.13)

    def card(self, slide, x, y, w, h, head, body, head_color=C.RED, fill=C.BG, num=None, head_size=14,
             body_size=11, bullet=True, head_h=None):
        self.box(slide, x, y, w, h, fill=fill)
        hx = x + 0.15
        if num is not None:
            self.num_badge(slide, x + 0.12, y + 0.12, num, color=head_color)
            hx = x + 0.56
        hh = head_h or 0.6
        self.text(slide, hx, y + 0.1, x + w - hx - 0.1, hh, head, size=head_size, color=head_color, bold=True,
                  anchor="t")
        self.text(slide, x + 0.15, y + 0.1 + hh, w - 0.3, h - hh - 0.2, body, size=body_size, bullet=bullet)

    def kpi(self, slide, x, y, w, h, value, label, sub=None, color=C.RED, fill=C.BG, vsize=24):
        self.box(slide, x, y, w, h, fill=fill)
        self.text(slide, x + 0.1, y + 0.06, w - 0.2, 0.3, label, size=10, color=C.SUB, align="c")
        self.text(slide, x + 0.05, y + 0.32, w - 0.1, h * 0.45, value, size=vsize, color=color, bold=True,
                  align="c", anchor="m")
        if sub:
            self.text(slide, x + 0.1, y + h - 0.34, w - 0.2, 0.3, sub, size=10, color=C.SUB, align="c")

    def callout(self, slide, x, y, w, h, label, body, color=C.RED):
        """スライド下部の「ひとこと」帯（淡色の全面ボックス。縁取り線は使わない）"""
        tint = C.RED_T if color == C.RED else C.NAVY_T
        self.box(slide, x, y, w, h, fill=tint)
        self.tag(slide, x + 0.12, y + (h - 0.28) / 2, 1.1, label, color=color, h=0.28)
        self.text(slide, x + 1.35, y, w - 1.45, h, body, size=12, anchor="m", bold=True, color=C.TEXT)

    def flow(self, slide, x, y, w, h, nodes, gap=0.32, colors=None, size=11, fills=None):
        """左→右の因果フロー。nodes: [str | (title, sub)]"""
        n = len(nodes)
        bw = (w - gap * (n - 1)) / n
        for i, nd in enumerate(nodes):
            bx = x + i * (bw + gap)
            col = (colors or [C.NAVY] * n)[i]
            fill = (fills or [None] * n)[i] or (C.RED_T if col == C.RED else C.NAVY_T if col == C.NAVY else C.BG)
            if isinstance(nd, tuple):
                paras = [[(nd[0], {"bold": True, "color": col, "size": size + 1})], [(nd[1], {"size": size - 1})]]
            else:
                paras = [[(nd, {"bold": True, "color": col})]]
            self.box(slide, bx, y, bw, h, fill=fill, text=paras, size=size)
            if i < n - 1:
                self.arrow(slide, bx + bw + 0.04, y + h / 2, bx + bw + gap - 0.04, y + h / 2, color=C.GRAY)

    def bar_chart(self, slide, x, y, w, h, cats, series, colors=None, title=None, number_format='0.0',
                  horizontal=False, legend=False, val_max=None, label_size=11, stacked=False, val_min=None,
                  label_colors=None):
        cd = CategoryChartData()
        cd.categories = cats
        for name, vals in series:
            cd.add_series(name, vals)
        if stacked:
            ct = XL_CHART_TYPE.BAR_STACKED if horizontal else XL_CHART_TYPE.COLUMN_STACKED
        else:
            ct = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
        gf = slide.shapes.add_chart(ct, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        ch = gf.chart
        ch.font.name = FONT
        ch.font.size = Pt(10)
        ch.has_legend = legend
        if title:
            ch.has_title = True
            ch.chart_title.text_frame.text = title
            for r in ch.chart_title.text_frame.paragraphs[0].runs:
                self._font(r, 11, C.SUB, bold=True)
        else:
            ch.has_title = False
        plot = ch.plots[0]
        plot.gap_width = 60
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = number_format
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.CENTER if stacked else XL_LABEL_POSITION.OUTSIDE_END  # 積み上げでOUTSIDE_ENDは破損
        dl.font.size = Pt(label_size)
        dl.font.bold = True
        dl.font.name = FONT
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = _rgb("E0E0E0")
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(9)
        va.tick_labels.font.color.rgb = _rgb(C.GRAY)
        if val_max is not None:
            va.maximum_scale = val_max
            va.minimum_scale = 0 if val_min is None else val_min
        elif val_min is not None:
            va.minimum_scale = val_min
        if legend:
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(10)
            from pptx.enum.chart import XL_LEGEND_POSITION
            ch.legend.position = XL_LEGEND_POSITION.TOP
        ca = ch.category_axis
        ca.tick_labels.font.size = Pt(10)
        ca.tick_labels.font.color.rgb = _rgb(C.SUB)
        ca.format.line.color.rgb = _rgb(C.LINE)
        for i, s in enumerate(plot.series):
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = _rgb((colors or [C.RED])[i % len(colors or [C.RED])])
            if label_colors:  # 系列ごとのラベル設定（表示・書式・位置も系列側で指定し直す）
                sdl = s.data_labels
                sdl.show_value = True
                sdl.number_format = number_format
                sdl.number_format_is_linked = False
                sdl.position = XL_LABEL_POSITION.CENTER if stacked else XL_LABEL_POSITION.OUTSIDE_END
                sdl.font.name = FONT
                sdl.font.color.rgb = _rgb(label_colors[i % len(label_colors)])
                s.data_labels.font.size = Pt(label_size)
                s.data_labels.font.bold = True
        return ch

    def line_chart(self, slide, x, y, w, h, cats, series, colors=None, title=None, number_format='0.0',
                   legend=False, val_min=None, val_max=None, label_size=10, dashed=None, labels=True):
        """折れ線グラフ（マーカー付き）。dashed: 破線にする系列番号のリスト（予想値など）"""
        from pptx.enum.chart import XL_MARKER_STYLE, XL_LEGEND_POSITION
        cd = CategoryChartData()
        cd.categories = cats
        for name, vals in series:
            cd.add_series(name, vals)
        gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        ch = gf.chart
        ch.font.name = FONT
        ch.font.size = Pt(10)
        ch.has_legend = legend
        if legend:
            ch.legend.include_in_layout = False
            ch.legend.position = XL_LEGEND_POSITION.TOP
            ch.legend.font.size = Pt(10)
        if title:
            ch.has_title = True
            ch.chart_title.text_frame.text = title
            for r in ch.chart_title.text_frame.paragraphs[0].runs:
                self._font(r, 11, C.SUB, bold=True)
        else:
            ch.has_title = False
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = _rgb("E0E0E0")
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(9)
        va.tick_labels.font.color.rgb = _rgb(C.GRAY)
        if val_min is not None:
            va.minimum_scale = val_min
        if val_max is not None:
            va.maximum_scale = val_max
        ca = ch.category_axis
        ca.tick_labels.font.size = Pt(10)
        ca.tick_labels.font.color.rgb = _rgb(C.SUB)
        ca.format.line.color.rgb = _rgb(C.LINE)
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        for i, s in enumerate(ch.plots[0].series):
            col = (colors or [C.RED])[i % len(colors or [C.RED])]
            s.format.line.color.rgb = _rgb(col)
            s.format.line.width = Pt(2.25)
            s.smooth = False
            s.marker.style = XL_MARKER_STYLE.CIRCLE
            s.marker.size = 7
            s.marker.format.fill.solid()
            s.marker.format.fill.fore_color.rgb = _rgb(col)
            s.marker.format.line.color.rgb = _rgb(col)
            if dashed and i in dashed:
                s.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            if labels:
                dl = s.data_labels
                dl.show_value = True
                dl.number_format = number_format
                dl.number_format_is_linked = False
                dl.position = XL_LABEL_POSITION.ABOVE
                dl.font.size = Pt(label_size)
                dl.font.bold = True
                dl.font.name = FONT
                dl.font.color.rgb = _rgb(col)
        return ch

    def table(self, slide, x, y, w, rows, col_w, row_h=0.32, header_fill=C.NAVY, size=10, header_size=10,
              zebra=True, aligns=None):
        nr, nc = len(rows), len(rows[0])
        gf = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
        tbl = gf.table
        # 既定スタイルを無効化
        tblPr = tbl._tbl.tblPr
        tblPr.set("bandRow", "0")
        tblPr.set("firstRow", "0")
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
        for r in range(nr):
            tbl.rows[r].height = Inches(row_h)
            for c in range(nc):
                cell = tbl.cell(r, c)
                val = rows[r][c]
                if val == "" or val is None:
                    val = " "  # 空セルは既定18ptになり行が高くなるため、空白1文字で文字サイズを効かせる
                cell.margin_left = cell.margin_right = Inches(0.06)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                head = r == 0
                al = "c" if head else ((aligns or ["l"] * nc)[c])
                paras = val if isinstance(val, list) else [val]
                tf.paragraphs[0].text = ""
                self._fill_tf(tf, paras, header_size if head else size, C.WHITE if head else C.TEXT, head,
                              al, False, 0)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(header_fill if head else (C.BG if zebra and r % 2 == 0 else C.WHITE))
        return tbl

    def view_matrix(self, slide, x, y, w, rows, comment_w=4.1, row_h=0.3, name_w=1.75):
        """CIO VIEW の5段階ドットマトリクス。
        rows: [dict(kind='group'|'asset', name, now(1-5), prev(1-5 or None), note)]"""
        scale_w = w - name_w - comment_w
        cw = scale_w / 5
        # 見出し
        for i, lab in enumerate(C.SCALE_LABEL):
            self.text(slide, x + name_w + i * cw, y, cw, 0.42, lab, size=9, color=C.SCALE[i], bold=True,
                      align="c", anchor="b", space_after=0, line_spacing=0.9)
        self.text(slide, x + name_w + scale_w + 0.1, y, comment_w - 0.1, 0.42, "ひとことコメント", size=9,
                  color=C.SUB, bold=True, anchor="b")
        cy = y + 0.48
        for rw in rows:
            grp = rw["kind"] == "group"
            fill = C.NAVY_T if grp else None
            if grp:
                self.box(slide, x, cy, w, row_h, fill=C.BG, shape="rect")
            self.text(slide, x + 0.05, cy, name_w - 0.1, row_h, rw["name"], size=11 if grp else 10.5,
                      bold=grp, anchor="m", color=C.TEXT)
            # 目盛り
            for i in range(5):
                cx = x + name_w + i * cw + cw / 2
                self.box(slide, cx - 0.03, cy + row_h / 2 - 0.03, 0.06, 0.06, fill=C.LINE, shape="oval")
            prev, now = rw.get("prev"), rw["now"]
            d = row_h * 0.72
            if prev and prev != now:
                px = x + name_w + (prev - 1) * cw + cw / 2
                nx = x + name_w + (now - 1) * cw + cw / 2
                self.box(slide, px - d / 2, cy + (row_h - d) / 2, d, d, fill=C.WHITE, line=C.SCALE[prev - 1],
                         shape="oval", dash=True, line_w=1.25)
                col = C.RED if now > prev else C.NAVY
                sgn = 1 if nx > px else -1
                self.arrow(slide, px + sgn * d / 2, cy + row_h / 2, nx - sgn * d / 2, cy + row_h / 2, color=col,
                           width=1.5)
            nx = x + name_w + (now - 1) * cw + cw / 2
            self.box(slide, nx - d / 2, cy + (row_h - d) / 2, d, d, fill=C.SCALE[now - 1], shape="oval")
            if rw.get("note"):
                self.text(slide, x + name_w + scale_w + 0.1, cy, comment_w - 0.1, row_h, rw["note"], size=9.5,
                          anchor="m", color=C.TEXT if not grp else C.SUB, bold=grp)
            cy += row_h + 0.02
        return cy

    def save(self, path):
        self.prs.save(path)
